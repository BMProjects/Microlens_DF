from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style:
        new_para.style = style
    return new_para


def insert_table_after(paragraph: Paragraph, rows: int, cols: int, style: str = "Table Grid") -> Table:
    table = paragraph._parent.add_table(rows=rows, cols=cols, width=Cm(16.5))
    table.style = style
    paragraph._p.addnext(table._tbl)
    return table


def set_cell_text(cell, text: str, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def find_paragraph(doc: Document, text: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip() == text:
            return para
    raise ValueError(f"Paragraph not found: {text}")


def pick_example_stems() -> list[str]:
    return ["123l_y0560_x1120", "100r_y0000_x1680"]


def build_example_pair(raw_path: Path, overlay_path: Path, title: str, out_path: Path) -> None:
    raw = Image.open(raw_path).convert("RGB")
    overlay = Image.open(overlay_path).convert("RGB")
    font = ImageFont.load_default()

    margin = 24
    title_h = 44
    label_h = 28
    pair_w = raw.width + overlay.width + margin * 3
    pair_h = title_h + label_h + raw.height + margin * 2

    canvas = Image.new("RGB", (pair_w, pair_h), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((margin, 12), title, fill="black", font=font)

    raw_x = margin
    img_y = title_h + label_h
    overlay_x = raw_x + raw.width + margin

    draw.text((raw_x, title_h), "原始测试图像", fill="black", font=font)
    draw.text((overlay_x, title_h), "对应标注示意图", fill="black", font=font)
    canvas.paste(raw, (raw_x, img_y))
    canvas.paste(overlay, (overlay_x, img_y))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(out_path, quality=95)


def create_assets(asset_dir: Path) -> list[Path]:
    img_dir = Path("output/tile_dataset/images/val")
    overlay_dir = Path("output/tile_dataset/overlays")
    figures: list[Path] = []
    for idx, stem in enumerate(pick_example_stems(), start=1):
        raw_path = img_dir / f"{stem}.jpg"
        overlay_path = overlay_dir / f"{stem}_overlay.jpg"
        figure_path = asset_dir / f"dataset_example_{idx}.jpg"
        build_example_pair(raw_path, overlay_path, f"示意样本 {idx}", figure_path)
        figures.append(figure_path)
    return figures


def update_reference_section(doc: Document) -> None:
    basis_para = find_paragraph(doc, "本大纲主要依据以下文件编制：")
    current = basis_para
    refs = [
        "GB/T 25000.51-2016《系统与软件工程 系统与软件质量要求与评价（SQuaRE） 第 51 部分：就绪可用软件产品（RUSP）的质量要求和测试细则》",
        "《青少年近视防控功能镜片关键参数计量技术研究》项目任务书（项目编号：2024YFC2419500）",
        "《离焦微结构镜片磨损程度检测及关键参数评估研究》课题任务书（课题编号：2024YFC2419504）",
        "委托测试文件、测试实施细则及经确认冻结的测试数据清单（如有）。",
    ]
    # clear any existing list items right after heading intro if they match old content
    for para in doc.paragraphs:
        if para.text.strip().startswith("`GB/T 25000.51-2016"):
            para.text = ""
    for ref in refs:
        current = insert_paragraph_after(current, ref, style="List Paragraph")


def update_flow_text(doc: Document) -> None:
    replacements = {
        "本大纲适用于对测试图像中的缺陷识别结果进行准确率测试。无论被测系统采用何种技术实现方式，只要能够对测试图中的缺陷进行识别并输出结果，均可依据本大纲开展测试。":
            "本大纲适用于对测试图像中磨损目标识别结果开展准确率测试。无论被测系统采用何种技术实现方式，只要能够对测试图像中的磨损目标进行识别并输出结果，均可依据本大纲组织测试。",
        "本大纲规定统一的测试原则、统计要求和结果形成方式，不对不同技术路线或不同输出形式分别设立独立指标体系。本大纲不规定具体达标数值，具体判定阈值应由任务书、委托测试文件或测试实施细则确定。":
            "本大纲规定统一的测试原则、统计要求、记录要求和结果形成方式，不预先限定具体图像数量、样本分布和磨损类别设置，以保证后续在数据扩充、类别扩展和技术路线调整时仍可延续使用。具体判定阈值由委托测试文件或相关考核要求确定。",
        "测试机构负责依据本大纲及委托测试文件实施测试，记录测试过程，核验测试结果，并编制测试报告":
            "测试机构负责依据本大纲及委托测试文件实施测试，完整记录测试过程，核验测试结果，并编制测试报告。",
        "测试机构应确保以下信息完整记录并可追溯：被测软件名称及版本、模型或算法版本、配置文件版本、测试数据集版本、基准结果版本、测试数据集及基准结果清单、运行环境信息、测试执行命令或运行方式、正式测试结果、重复性核验结果、原始输出结果文件、统计结果文件、测试日志以及异常情况记录。":
            "测试机构应确保以下信息完整记录并可追溯：被测软件名称及版本、模型或算法版本、配置文件版本、测试数据集版本、基准结果版本、测试数据集及基准结果清单、运行环境信息、测试执行命令或运行方式、正式测试结果、重复性核验结果、原始输出结果文件、统计结果文件、测试日志以及异常情况记录。必要时可结合附录中的记录表格进行登记。",
        "本大纲覆盖的测试项目为“磨损识别准确率测试”。":
            "本大纲覆盖的测试项目为“磨损识别准确率测试”，测试过程中仅围绕该指标开展，不扩展至其他性能评价项。",
        "测试内容为：在规定测试图像集上，对图像中的镜片磨损进行识别，并将被测系统的识别结果与人工确认的基准结果进行比较，得到识别准确率。":
            "测试内容为：在规定测试图像集上，由被测系统对图像中的镜片磨损目标进行识别，并将识别结果与人工确认并冻结的基准结果进行比较，得到磨损识别准确率。",
        "测试项目列表说明如下：":
            "测试项目列表说明如下。为便于说明测试数据的组织形式和标注形式，以下给出两组简单、易于识别的示意样本。示意样本仅用于说明数据形态，不构成对正式测试图像数量、样本分布或磨损类别设置的限定。",
        "测试时应满足以下要求：":
            "正式测试时应满足以下要求：",
        "每次正式测试前，应明确并登记以下测试数据要素：":
            "每次正式测试前，应对以下测试数据要素进行确认和登记：",
        "测试环境应满足被测软件完整运行的要求，并具备可重复执行测试的条件。":
            "测试环境应满足被测软件完整运行的要求，并具备稳定、可复现和可重复执行测试的条件。",
        "被测软件应提供软件名称和版本标识、模型或算法版本标识、运行命令或启动方式、输入输出说明以及必要的配置文件。":
            "被测软件应提供软件名称和版本标识、模型或算法版本标识、运行命令或启动方式、输入输出说明以及必要的配置文件，并确保第三方测试机构能够据此独立完成测试实施。",
        "测试步骤如下：":
            "测试步骤如下。测试过程中的数据冻结、测试执行和异常样本处理，可结合附录中的记录表格同步留痕：",
        "每次测试至少应形成以下输出物：测试数据集清单、正式测试日志、重复性核验日志、原始识别输出结果、指标统计结果文件、测试记录表和正式测试报告。":
            "每次测试至少应形成以下输出物：测试数据集清单、正式测试日志、重复性核验日志、原始识别输出结果、指标统计结果文件、测试记录表和正式测试报告。附录A至附录C可作为测试过程的统一记录模板。",
    }
    for para in doc.paragraphs:
        src = para.text.strip()
        if src in replacements:
            para.text = replacements[src]


def update_metric_section(doc: Document) -> None:
    mapping = {
        72: "磨损识别准确率采用统一指标名称、统一判定规则和统一统计口径进行计算。",
        73: "设纳入本次统计的测试样本总数为 N。对于第 i 个测试样本，记被测软件输出结果为 R_i，对应基准结果为 G_i。依据委托测试文件或测试实施细则中规定的一致性判定规则，对 R_i 与 G_i 进行比较，并记该样本的一致性判定结果为 δ_i。",
        74: "当样本识别结果满足既定一致性判定规则时，可记 δ_i = 1；当样本识别结果不满足既定一致性判定规则时，可记 δ_i = 0；对于采用连续一致性评分的情况，δ_i 亦可取 0 至 1 之间的实数。",
        75: "则总体磨损识别准确率 A 按下式计算：",
        76: "A = (1/N) × Σ(i=1→N) δ_i × 100%",
        77: "其中：",
        78: "N 为纳入本次测试统计的测试样本总数；",
        79: "R_i 为第 i 个测试样本的被测软件输出结果；",
        80: "G_i 为第 i 个测试样本的基准结果；",
        81: "δ_i 为第 i 个测试样本的一致性判定结果；",
        82: "A 为总体磨损识别准确率。",
        83: "一致性判定规则、样本纳入原则、异常样本处理方式以及合格判定阈值，应由任务书、委托测试文件或测试实施细则明确，并在同一次测试中保持一致。",
    }
    for idx, text in mapping.items():
        doc.paragraphs[idx].text = text
        if idx == 76:
            doc.paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.CENTER


def insert_dataset_examples(doc: Document, figures: list[Path]) -> None:
    anchor = find_paragraph(
        doc,
        "测试项目列表说明如下。为便于说明测试数据的组织形式和标注形式，以下给出两组简单、易于识别的示意样本。示意样本仅用于说明数据形态，不构成对正式测试图像数量、样本分布或磨损类别设置的限定。",
    )
    current = anchor
    note = insert_paragraph_after(
        current,
        "示意样本说明：每组示意图均由原始测试图像和对应标注示意图组成，左图为原始图像，右图为经人工标注后的示意结果。",
        style="Normal",
    )
    current = note
    for idx, fig in enumerate(figures, start=1):
        p = insert_paragraph_after(current, style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(fig), width=Cm(16.5))
        caption = insert_paragraph_after(
            p,
            f"图 {idx} 测试图像与对应标注示意图（示意样本）",
            style="Normal",
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        current = caption


def append_appendices(doc: Document) -> None:
    current = doc.paragraphs[-1]

    current = insert_paragraph_after(current, "附录A 测试数据冻结登记表（示例）", style="Heading 1")
    current = insert_paragraph_after(
        current,
        "附录A用于记录正式测试前所采用测试数据集、基准结果及样本清单的冻结信息，可根据实际测试需要增减行数。",
        style="Normal",
    )
    table_a = insert_table_after(current, rows=2, cols=8)
    table_a.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_a = [
        "测试数据集版本号",
        "基准结果版本号",
        "样本清单文件路径",
        "原图数量",
        "纳入统计样本数量",
        "基准目标总数",
        "冻结确认日期",
        "备注",
    ]
    for i, h in enumerate(headers_a):
        set_cell_text(table_a.rows[0].cells[i], h, bold=True, center=True)
    for i in range(len(headers_a)):
        set_cell_text(table_a.rows[1].cells[i], " ", center=False)

    current = insert_paragraph_after(current, "附录B 测试执行与重复性核验记录表（示例）", style="Heading 1")
    current = insert_paragraph_after(
        current,
        "附录B用于记录正式测试和重复性核验测试的执行条件与结果，便于后续核查测试过程是否保持一致。",
        style="Normal",
    )
    table_b = insert_table_after(current, rows=3, cols=8)
    table_b.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_b = [
        "测试阶段",
        "软件版本",
        "模型/权重版本",
        "配置文件版本",
        "测试命令",
        "输出目录",
        "主指标结果",
        "备注",
    ]
    for i, h in enumerate(headers_b):
        set_cell_text(table_b.rows[0].cells[i], h, bold=True, center=True)
    set_cell_text(table_b.rows[1].cells[0], "正式测试")
    set_cell_text(table_b.rows[2].cells[0], "重复性核验")
    for row in table_b.rows[1:]:
        for cell in row.cells[1:]:
            set_cell_text(cell, " ")

    current = insert_paragraph_after(current, "附录C 异常样本记录表（示例）", style="Heading 1")
    current = insert_paragraph_after(
        current,
        "附录C用于记录测试过程中发现的异常样本及其处理情况，以保证正式统计过程中的样本处理口径一致。",
        style="Normal",
    )
    table_c = insert_table_after(current, rows=2, cols=6)
    table_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_c = ["序号", "样本标识", "异常类型", "是否纳入统计", "处理说明", "记录日期"]
    for i, h in enumerate(headers_c):
        set_cell_text(table_c.rows[0].cells[i], h, bold=True, center=True)
    for i in range(len(headers_c)):
        set_cell_text(table_c.rows[1].cells[i], " ")


def update_table(doc: Document) -> None:
    table = doc.tables[0]
    set_cell_text(table.rows[1].cells[3], "磨损识别准确率（按委托测试文件明确的评价方式计算）")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--asset-dir", default="cnas_test/docs/assets/outline_examples")
    parser.add_argument("--sync-current", action="store_true")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    asset_dir = Path(args.asset_dir)

    figures = create_assets(asset_dir)
    doc = Document(str(input_path))

    update_reference_section(doc)
    update_flow_text(doc)
    update_metric_section(doc)
    update_table(doc)
    insert_dataset_examples(doc, figures)
    append_appendices(doc)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    if args.sync_current:
        current_path = output_path.parent / "CNAS测试大纲_当前版.docx"
        shutil.copyfile(output_path, current_path)


if __name__ == "__main__":
    main()
