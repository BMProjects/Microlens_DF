from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.text.paragraph import Paragraph


def set_paragraph(
    paragraph: Paragraph,
    text: str,
    *,
    style: str | None = None,
    align: int | None = None,
) -> None:
    paragraph.text = ""
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    if align is not None:
        paragraph.alignment = align


def clear_paragraph(paragraph: Paragraph) -> None:
    paragraph.text = ""
    paragraph.style = "Normal"


def set_picture(paragraph: Paragraph, image_path: Path, width_cm: float = 7.5) -> None:
    paragraph.text = ""
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True)
    args = parser.parse_args()

    path = Path(args.docx)
    doc = Document(path)
    asset_dir = Path("cnas_test/docs/assets/outline_examples_full")
    fig1 = asset_dir / "full_image_example_1.png"
    fig2 = asset_dir / "full_image_example_2.png"

    # Section 4
    set_paragraph(doc.paragraphs[37], "4 测试组织管理和流程", style="Heading 1")
    set_paragraph(doc.paragraphs[38], "4.1 测试组织管理", style="Heading 2")
    set_paragraph(
        doc.paragraphs[39],
        "测试组织管理应符合第三方软件测试的一般要求，并体现独立、客观、可追溯和可重复执行的原则。测试前应确认送测资料、测试范围和结果判定要求，测试过程中应保留必要记录。",
        style="Normal",
    )
    set_paragraph(doc.paragraphs[40], "测试组织和职责分工如下：", style="Normal")
    set_paragraph(
        doc.paragraphs[41],
        "委托方负责提出测试需求，并提交智能磨损识别算法的送测版本、运行说明、必要配置文件、测试图像集、基准结果以及磨损镜片图像清单等资料。",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[42],
        "测试机构负责依据本大纲实施测试，完成环境核验、测试执行、结果核验、必要复核及测试报告编制。",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[43],
        "测试过程中形成的运行记录、输出结果、统计结果和异常记录应统一保存，作为结果核验、测试结论形成和报告编制的依据。测试机构应保证测试资料、运行条件、原始结果和统计结果之间能够相互对应、便于复核。",
        style="Normal",
    )
    set_paragraph(doc.paragraphs[44], "4.2 测试流程", style="Heading 2")
    set_paragraph(
        doc.paragraphs[45],
        "测试一般按照“资料确认—环境核验—正式测试—结果核验—报告形成”的顺序开展。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[46],
        "接收并确认委托测试要求，以及本次送测智能磨损识别算法的名称、版本和交付资料；",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[47],
        "核对测试图像集、基准结果、磨损镜片图像清单和必要配置文件，完成测试准备；",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[48],
        "核验测试环境和运行条件，确认图像读取、结果输出和日志记录满足测试要求；",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[49],
        "按规定环境和测试图像集执行正式测试，生成原始结果、统计结果和日志；",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[50],
        "对正式测试结果进行核验，并在需要时开展重复性核验；",
        style="List Paragraph",
    )
    set_paragraph(
        doc.paragraphs[51],
        "汇总测试记录和核验结果，形成测试结论，并编制测试报告。",
        style="Normal",
    )
    clear_paragraph(doc.paragraphs[52])
    clear_paragraph(doc.paragraphs[53])

    # Section 5
    set_paragraph(doc.paragraphs[54], "5 测试用例和测试对象", style="Heading 1")
    set_paragraph(doc.paragraphs[55], "5.1 磨损镜片图像", style="Heading 2")
    set_paragraph(
        doc.paragraphs[56],
        "本大纲以纳入测试图像集的单张磨损镜片图像作为基本测试单元。正式测试使用的磨损镜片图像应来自经确认的测试图像集，并与对应基准结果保持一一对应。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[57],
        "每张磨损镜片图像均应具有唯一图像标识，便于测试机构按图像逐项核对识别结果、统计结果和测试记录。下列示例图像仅用于说明测试图像的基本形态。",
        style="Normal",
    )
    set_picture(doc.paragraphs[58], fig1)
    set_paragraph(doc.paragraphs[59], "图 1 磨损镜片图像示例", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_picture(doc.paragraphs[60], fig2)
    set_paragraph(doc.paragraphs[61], "图 2 磨损镜片图像示例", style="Normal", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph(
        doc.paragraphs[62],
        "测试准备时，宜形成磨损镜片图像清单，记录图像标识、文件路径、对应基准结果版本及必要备注信息。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[63],
        "正式测试前，应完成测试图像集版本、基准结果版本和磨损镜片图像清单的一致性核对。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[64],
        "测试期间使用的磨损镜片图像及其基准结果应保持稳定，以保证测试过程和测试结果可以复核。",
        style="Normal",
    )
    set_paragraph(doc.paragraphs[65], "5.2 智能磨损识别算法", style="Heading 2")
    set_paragraph(
        doc.paragraphs[66],
        "本次接受测试的对象为智能磨损识别算法。算法应能够读取规定格式的磨损镜片图像，并输出可用于统计和核验的识别结果。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[67],
        "提交测试时，宜同时提供算法版本、模型或参数版本、运行命令或启动方式、必要配置文件、输入输出说明以及结果保存方式。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[68],
        "对于批量处理方式，算法输出结果应与磨损镜片图像标识保持一致，便于测试机构逐图像核验和统计。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[69],
        "算法运行过程中形成的原始结果、统计结果和必要日志应能够留存，以满足测试核验和结果追溯的需要。",
        style="Normal",
    )

    # Section 6
    set_paragraph(doc.paragraphs[70], "6 性能指标测试", style="Heading 1")
    set_paragraph(doc.paragraphs[71], "6.1 测试说明和要求", style="Heading 2")
    set_paragraph(
        doc.paragraphs[72],
        "本次测试项目为“磨损识别准确率测试”。测试时，以智能磨损识别算法对规定测试图像集的识别结果与基准结果的一致程度作为评价依据。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[73],
        "同一次测试应采用统一的磨损镜片图像、基准结果、统计规则和判定方式，并保持测试数据和测试条件前后一致。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[74],
        "如需进行重复性核验，应在相同环境、相同数据、相同算法版本和相同配置条件下进行。",
        style="Normal",
    )
    set_paragraph(doc.paragraphs[75], "6.2 测试环境及条件", style="Heading 2")
    set_paragraph(
        doc.paragraphs[76],
        "测试环境应满足智能磨损识别算法完整运行的要求，并具备稳定、可复现和可重复执行测试的条件。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[77],
        "正式测试前，应确认运行环境、输入输出路径、配置文件和相关访问权限满足测试要求，并保证测试过程中形成的原始结果、统计结果和日志能够完整保存、便于核验和追溯。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[78],
        "对于批量处理方式，测试环境应支持按磨损镜片图像清单顺序稳定运行，并保持图像输入、结果输出和日志记录之间的一致对应。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[79],
        "正式测试与重复性核验应在一致的运行条件下完成。测试过程中如发生影响结果形成的环境变化、配置变化或资料变更，应停止当前测试并重新确认测试条件。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[80],
        "测试过程中形成的日志、原始结果和统计结果应妥善保存，以便核验和追溯。",
        style="Normal",
    )
    clear_paragraph(doc.paragraphs[81])
    clear_paragraph(doc.paragraphs[82])
    set_paragraph(doc.paragraphs[83], "6.3 测试方法和步骤", style="Heading 2")
    set_paragraph(
        doc.paragraphs[84],
        "磨损识别准确率按统一统计口径计算。设纳入本次统计的磨损镜片图像总数为 N。对第 i 个磨损镜片图像，记算法输出结果为 R_i，对应基准结果为 G_i。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[85],
        "依据本次测试确认的判定规则，对 R_i 与 G_i 进行比较，并记该图像的一致性判定结果为 δ_i。",
        style="Normal",
    )
    set_paragraph(doc.paragraphs[86], "δ_i = 1（识别结果满足本次测试确认的一致性判定规则）", style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph(doc.paragraphs[87], "则总体磨损识别准确率 A 按下式计算：", style="Normal")
    set_paragraph(doc.paragraphs[88], "A = (1/N) × Σ(i=1→N) δ_i × 100%", style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph(
        doc.paragraphs[89],
        "式中，N 为纳入统计的磨损镜片图像总数；R_i 为第 i 个磨损镜片图像的算法输出结果；G_i 为对应基准结果；δ_i 为该图像的一致性判定结果；A 为总体磨损识别准确率。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[90],
        "测试实施时，测试机构应先准备并核对测试图像集、基准结果和磨损镜片图像清单，再按登记确认的运行条件执行正式测试，保存原始结果、统计结果和日志文件。",
        style="Normal",
    )
    set_paragraph(
        doc.paragraphs[91],
        "正式测试完成后，应按统一规则进行逐图像核验，并视需要开展重复性核验，最终形成测试结论和测试报告。",
        style="Normal",
    )

    doc.save(path)


if __name__ == "__main__":
    main()
