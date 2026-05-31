from __future__ import annotations

import argparse
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, ascii_font: str = "Times New Roman", east_asia_font: str = "宋体", size: int = 12) -> None:
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)


def add_inline(paragraph, node, *, default_ascii: str = "Times New Roman", default_east_asia: str = "宋体", size: int = 12, bold: bool = False, italic: bool = False) -> None:
    if isinstance(node, NavigableString):
        text = str(node)
        if not text:
            return
        run = paragraph.add_run(text)
        set_run_font(run, default_ascii, default_east_asia, size)
        run.bold = bold
        run.italic = italic
        return

    if not isinstance(node, Tag):
        return

    tag = node.name.lower()
    next_bold = bold or tag in {"strong", "b"}
    next_italic = italic or tag in {"em", "i"}

    if tag == "br":
        paragraph.add_run().add_break()
        return

    if tag == "code":
        for child in node.children:
            add_inline(
                paragraph,
                child,
                default_ascii=default_ascii,
                default_east_asia=default_east_asia,
                size=size,
                bold=next_bold,
                italic=next_italic,
            )
        return

    if tag == "sub":
        text = f"_{node.get_text()}"
        run = paragraph.add_run(text)
        set_run_font(run, default_ascii, default_east_asia, size)
        run.bold = next_bold
        run.italic = next_italic
        return

    if tag == "sup":
        text = f"^{node.get_text()}"
        run = paragraph.add_run(text)
        set_run_font(run, default_ascii, default_east_asia, size)
        run.bold = next_bold
        run.italic = next_italic
        return

    for child in node.children:
        add_inline(
            paragraph,
            child,
            default_ascii=default_ascii,
            default_east_asia=default_east_asia,
            size=size,
            bold=next_bold,
            italic=next_italic,
        )


def add_paragraph_from_tag(doc: Document, tag: Tag, *, style: str | None = None, align=None, size: int = 12, space_after: int = 6) -> None:
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    for child in tag.children:
        add_inline(paragraph, child, size=size)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for style_name, size, east_asia in [
        ("Heading 1", 16, "黑体"),
        ("Heading 2", 14, "黑体"),
        ("Heading 3", 12, "黑体"),
        ("Heading 4", 12, "黑体"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
        style.font.size = Pt(size)
        style.font.bold = True

    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 1"].paragraph_format.line_spacing = 1.2
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.line_spacing = 1.2
    doc.styles["Heading 3"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 3"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 3"].paragraph_format.line_spacing = 1.2
    doc.styles["Heading 4"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 4"].paragraph_format.space_after = Pt(3)
    doc.styles["Heading 4"].paragraph_format.line_spacing = 1.2

    for list_style_name in ["List Bullet", "List Number"]:
        style = doc.styles[list_style_name]
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.25


def render_cover(doc: Document, cover: Tag) -> None:
    for div in cover.find_all("div", recursive=False):
        cls = div.get("class", [])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "title" in cls:
            run = p.add_run(div.get_text(strip=True))
            set_run_font(run, "Times New Roman", "黑体", 24)
            run.bold = True
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(18)
        elif "proj" in cls:
            run = p.add_run(div.get_text(strip=True))
            set_run_font(run, "Times New Roman", "宋体", 12)
        elif "kv" in cls:
            for line in div.find_all("div", recursive=False):
                lp = doc.add_paragraph()
                lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = lp.add_run(line.get_text(" ", strip=True).replace("\xa0", ""))
                set_run_font(run, "Times New Roman", "宋体", 12)
    doc.add_page_break()


def render_table(doc: Document, table_tag: Tag) -> None:
    rows = table_tag.find_all("tr", recursive=False)
    if not rows:
        return
    first_cells = rows[0].find_all(["th", "td"], recursive=False)
    table = doc.add_table(rows=0, cols=len(first_cells))
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row_tag in enumerate(rows):
        cells = row_tag.find_all(["th", "td"], recursive=False)
        row = table.add_row()
        for c_idx, cell_tag in enumerate(cells):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            for child in cell_tag.children:
                add_inline(p, child, size=10 if r_idx == 0 else 10)
            if cell_tag.name.lower() == "th":
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(6)


def render_figure(doc: Document, fig_div: Tag, html_path: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    for img in fig_div.find_all("img", recursive=False):
        src = img.get("src", "")
        img_path = (html_path.parent / src).resolve()
        if img_path.exists():
            run = p.add_run()
            run.add_picture(str(img_path), width=Cm(5.9))
            p.add_run("  ")


def render_main(doc: Document, main_tag: Tag, html_path: Path) -> None:
    for child in main_tag.children:
        if isinstance(child, NavigableString):
            continue
        if not isinstance(child, Tag):
            continue
        if child.name == "nav":
            continue
        if child.name == "div" and "cover" in child.get("class", []):
            render_cover(doc, child)
            continue
        if child.name == "h1":
            add_paragraph_from_tag(doc, child, style="Heading 1", size=16, space_after=10)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            doc.paragraphs[-1].paragraph_format.keep_together = True
            continue
        if child.name == "h2":
            add_paragraph_from_tag(doc, child, style="Heading 1", size=15, space_after=6)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            doc.paragraphs[-1].paragraph_format.keep_together = True
            continue
        if child.name == "h3":
            add_paragraph_from_tag(doc, child, style="Heading 2", size=13, space_after=4)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            doc.paragraphs[-1].paragraph_format.keep_together = True
            continue
        if child.name == "h4":
            add_paragraph_from_tag(doc, child, style="Heading 3", size=12, space_after=4)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            doc.paragraphs[-1].paragraph_format.keep_together = True
            continue
        if child.name == "p":
            add_paragraph_from_tag(doc, child, size=12)
            continue
        if child.name == "ul":
            for li in child.find_all("li", recursive=False):
                add_paragraph_from_tag(doc, li, style="List Bullet", size=11, space_after=2)
            continue
        if child.name == "ol":
            for li in child.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(2)
                for li_child in li.contents:
                    if isinstance(li_child, Tag) and li_child.name == "ul":
                        continue
                    add_inline(p, li_child, size=11)
                sublists = [x for x in li.contents if isinstance(x, Tag) and x.name == "ul"]
                for sub_ul in sublists:
                    for sub_li in sub_ul.find_all("li", recursive=False):
                        add_paragraph_from_tag(doc, sub_li, style="List Bullet 2", size=10, space_after=1)
            continue
        if child.name == "table":
            render_table(doc, child)
            continue
        if child.name == "div" and "formula" in child.get("class", []):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for formula_child in child.children:
                add_inline(p, formula_child, default_ascii="Times New Roman", default_east_asia="宋体", size=12)
            continue
        if child.name == "div" and "fig" in child.get("class", []):
            render_figure(doc, child, html_path)
            continue
        if child.name == "div" and "figline" in child.get("class", []):
            add_paragraph_from_tag(doc, child, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=8)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            continue


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--html", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    html_path = Path(args.html)
    out_path = Path(args.output)
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    main_tag = soup.find("main", class_="page")
    if main_tag is None:
        raise SystemExit("未找到主文档区域 <main class='page'>")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.start_type = WD_SECTION_START.NEW_PAGE

    configure_styles(doc)
    render_main(doc, main_tag, html_path)
    doc.save(out_path)


if __name__ == "__main__":
    main()
