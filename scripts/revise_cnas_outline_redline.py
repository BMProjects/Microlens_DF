from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


RED = RGBColor(0xFF, 0x00, 0x00)


def set_paragraph_text(paragraph: Paragraph, text: str, *, red: bool = False, bold: bool = False,
                       align: int | None = None) -> None:
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    if red:
        run.font.color.rgb = RED
    if align is not None:
        paragraph.alignment = align


def add_red_run(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    run.font.color.rgb = RED


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None,
                           red: bool = True, align: int | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        set_paragraph_text(new_para, text, red=red, align=align)
    elif align is not None:
        new_para.alignment = align
    return new_para


def insert_table_after(paragraph: Paragraph, rows: int, cols: int, style: str = "Table Grid") -> Table:
    table = paragraph._parent.add_table(rows=rows, cols=cols, width=Cm(16.5))
    table.style = style
    paragraph._p.addnext(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    return table


def insert_paragraph_after_table(table: Table, text: str = "", style: str | None = None,
                                 red: bool = True, align: int | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    if style:
        new_para.style = style
    if text:
        set_paragraph_text(new_para, text, red=red, align=align)
    elif align is not None:
        new_para.alignment = align
    return new_para


def set_cell(cell, text: str, *, red: bool = False, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    if red:
        run.font.color.rgb = RED
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def create_figure_assets(asset_dir: Path) -> list[Path]:
    stems = ["15r", "92l"]
    src_dir = Path("output/dataset_v2/images")
    asset_dir.mkdir(parents=True, exist_ok=True)
    out_paths: list[Path] = []
    for idx, stem in enumerate(stems, start=1):
        src = src_dir / f"{stem}.png"
        dst = asset_dir / f"full_image_example_{idx}.png"
        shutil.copyfile(src, dst)
        out_paths.append(dst)
    return out_paths


def replace_first_paragraph_containing(doc: Document, needle: str, new_text: str) -> bool:
    for para in doc.paragraphs:
        if needle in para.text:
            set_paragraph_text(para, new_text, red=True)
            return True
    return False


def replace_in_all_paragraphs(doc: Document, old: str, new: str) -> None:
    for para in doc.paragraphs:
        if old in para.text:
            set_paragraph_text(para, para.text.replace(old, new), red=True)


def style_picture_caption(paragraph: Paragraph, text: str) -> Paragraph:
    set_paragraph_text(paragraph, text, red=True)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return paragraph


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sync-current", action="store_true")
    parser.add_argument("--asset-dir", default="cnas_test/docs/assets/outline_examples_full")
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    asset_dir = Path(args.asset_dir)

    doc = Document(str(input_path))
    figs = create_figure_assets(asset_dir)

    # body table
    body_table = doc.tables[0]
    set_cell(body_table.rows[1].cells[1], "验证智能磨损识别算法对测试图中磨损目标识别结果与基准结果的一致程度。", red=True)
    set_cell(body_table.rows[1].cells[2], "磨损目标识别结果与基准结果的对应比较与统计。", red=True)
    set_cell(body_table.rows[1].cells[3], "磨损识别准确率（按本次测试确认的统计方式计算）。", red=True)

    # paragraph replacements
    updates = {
        7: "本大纲适用于对测试图像中磨损目标识别结果进行准确率测试。智能磨损识别算法只要能够针对磨损镜片图像给出可与基准结果进行对应比较的识别结果，均可依据本大纲开展测试。",
        8: "本大纲规定统一的测试原则、统计要求和结果形成方式。测试所采用的图像数据、基准结果、统计规则和判定阈值，以测试实施阶段经确认的版本为准。",
        14: "GB/T 25000.51-2016《系统与软件工程 系统与软件质量要求与评价（SQuaRE） 第 51 部分：就绪可用软件产品（RUSP）的质量要求和测试细则》",
        15: "GB/T 41864-2022《信息技术 计算机视觉 术语》",
        20: "委托方负责提出测试需求，提交智能磨损识别算法的软件包或部署版本、运行说明、版本信息、必要配置文件、测试图像集、基准结果以及磨损镜片图像清单等测试资料。",
        21: "测试机构负责依据本大纲及委托测试要求，对智能磨损识别算法进行测试准备审查、环境核验、测试实施、结果核验、重复性核验和测试报告编制。",
        22: "测试机构应确保以下信息完整记录并可追溯：智能磨损识别算法名称及版本、模型或算法版本、配置文件版本、测试图像集版本、基准结果版本、磨损镜片图像清单、运行环境信息、测试执行命令或运行方式、试运行结果、正式测试结果、重复性核验结果、原始输出结果文件、统计结果文件、测试日志以及异常情况记录。",
        25: "接收委托测试要求，确认测试范围、测试依据、结果判定要求以及本次送测智能磨损识别算法的名称、版本和交付形态；",
        26: "核对智能磨损识别算法的软件包、模型或算法版本、配置文件、运行说明、测试图像集、基准结果和磨损镜片图像清单，并完成测试准备登记；",
        27: "核验测试环境及运行条件，检查算法启动、图像读取、结果输出、日志记录和文件保存功能是否正常，确认正式测试和重复性核验所用环境保持一致；",
        28: "在不纳入正式统计的前提下开展必要试运行，确认智能磨损识别算法能够按磨损镜片图像清单顺序处理图像，并输出与图像标识对应的识别结果；",
        29: "执行正式测试，对规定测试图像集中的磨损镜片图像完成批量识别处理，生成完整的原始输出结果、统计结果和测试日志；",
        30: "依据确认后的统计规则对正式测试结果进行逐图像核验，并检查输出文件、统计结果和异常记录的完整性；",
        31: "在相同环境、相同输入和相同版本条件下执行重复性核验，形成重复性核验结果，并核对正式测试与重复性核验的一致性；",
        32: "汇总测试记录、核验结果和异常处理记录，形成测试结论，并编制测试报告。",
        10: "测试目的是验证智能磨损识别算法在规定测试图像集上对镜片磨损目标的识别结果，与人工确认并冻结的基准结果之间的一致程度是否达到委托测试文件规定的要求，并形成可供委托方、测试机构及验收使用的客观测试结论。",
        35: "本大纲覆盖的测试项目为“磨损识别准确率测试”。本次接受测试的内容为能够接收磨损镜片图像、完成磨损目标识别并输出可统计结果的智能磨损识别算法。",
        36: "测试内容为：在规定测试图像集上，由智能磨损识别算法对图像中的镜片磨损目标进行识别，并将识别结果与人工确认并冻结的基准结果进行比较，得到磨损识别准确率。",
        37: "测试项目列表说明如下。为便于说明测试数据的组织形式，以下给出两张磨损镜片图像示例。示例仅用于说明测试图像形态，不作为正式测试图像规模、样本分布或磨损类别设置的限定。",
        38: "测试时应满足以下要求：",
        39: "测试图像应独立于训练、调参和模型选择过程；",
        40: "测试图像及其基准结果应在测试前完成确认并冻结；",
        41: "测试期间不得临时替换、删减、增补磨损镜片图像，不得擅自修改基准结果；",
        42: "同一次测试中，识别结果的统计口径、匹配规则和判定规则应保持一致。",
        44: "测试环境应满足智能磨损识别算法完整运行的要求，并具备稳定、可复现和可重复执行测试的条件。",
        51: "智能磨损识别算法应提供软件名称和版本标识、模型或算法版本标识、运行命令或启动方式、输入输出说明以及必要的配置文件，使测试机构能够依据文档完成测试实施。对于批量处理磨损镜片图像的情况，还应能按图像标识输出对应结果。",
        52: "测试环境应能完整记录测试日志、输出结果和异常信息，以保证测试过程可追溯，并支持正式测试与重复性核验的一致性核对。",
        54: "磨损识别准确率采用统一指标名称、统一判定规则和统一统计口径进行计算。",
        55: "设纳入本次统计的磨损镜片图像总数为 N。对于第 i 个磨损镜片图像，记智能磨损识别算法输出结果为 R_i，对应基准结果为 G_i。依据本次测试确认的统计规则，对 R_i 与 G_i 进行比较，并记该图像的一致性判定结果为 δ_i：",
        57: "则总体磨损识别准确率 A 按下式计算：",
        59: "其中：",
        60: "N 为纳入本次测试统计的磨损镜片图像总数；",
        61: "R_i 为第 i 个磨损镜片图像的智能磨损识别算法输出结果；",
        62: "G_i 为第 i 个磨损镜片图像的基准结果；",
        63: "δ_i 为第 i 个磨损镜片图像的一致性判定结果；",
        64: "A 为总体磨损识别准确率。",
        65: "统计规则、样本纳入原则、异常样本处理方式及合格判定阈值，以本次测试确认后的版本执行，并在同一次测试中保持一致。",
        66: "测试步骤如下：",
        67: "准备经确认的测试图像集、基准结果、磨损镜片图像清单及测试执行记录表，逐项核对图像标识、文件路径、基准结果引用方式和结果保存位置；",
        68: "在规定测试环境中按磨损镜片图像清单顺序导入图像，完成必要试运行后执行正式测试，获得全部磨损镜片图像的识别结果，并保存日志和原始结果文件；",
        69: "依据本次测试确认的统计规则，对正式测试结果与基准结果进行逐图像比较，记录每张磨损镜片图像的处理状态、识别结果对应情况和异常信息，完成指标统计和结果核验；",
        70: "在相同测试环境、相同输入数据、相同智能磨损识别算法版本和相同配置条件下执行 1 次重复性核验，并对正式测试所形成的关键结果进行复核；",
        71: "核对正式测试结果与重复性核验结果，形成一致性判断，并据此形成测试结论。",
    }
    for idx, text in updates.items():
        set_paragraph_text(doc.paragraphs[idx], text, red=True)

    # Early paragraphs in the source doc have shifted across versions; replace by content anchor.
    replace_first_paragraph_containing(
        doc,
        "本大纲适用于对测试图像中的缺陷识别结果进行准确率测试",
        "本大纲适用于对测试图像中磨损目标识别结果进行准确率测试。智能磨损识别算法只要能够针对磨损镜片图像给出可与基准结果进行对应比较的识别结果，均可依据本大纲开展测试。",
    ) or replace_first_paragraph_containing(
        doc,
        "本大纲适用于对测试图像中磨损目标识别结果进行准确率测试",
        "本大纲适用于对测试图像中磨损目标识别结果进行准确率测试。智能磨损识别算法只要能够针对磨损镜片图像给出可与基准结果进行对应比较的识别结果，均可依据本大纲开展测试。",
    )
    replace_first_paragraph_containing(
        doc,
        "本大纲规定统一的测试原则、统计要求和结果形成方式，不对不同技术路线或不同输出形式分别设立独立指标体系。",
        "本大纲规定统一的测试原则、统计要求和结果形成方式。测试所采用的图像数据、基准结果、统计规则和判定阈值，以测试实施阶段经确认的版本为准。",
    ) or replace_first_paragraph_containing(
        doc,
        "本大纲规定统一的测试原则、统计要求和结果形成方式",
        "本大纲规定统一的测试原则、统计要求和结果形成方式。测试所采用的图像数据、基准结果、统计规则和判定阈值，以测试实施阶段经确认的版本为准。",
    )
    replace_first_paragraph_containing(
        doc,
        "测试目的是验证被测软件产品在规定测试图像集上对镜片磨损目标的识别结果",
        "测试目的是验证智能磨损识别算法在规定测试图像集上对镜片磨损目标的识别结果，与人工确认并冻结的基准结果之间的一致程度是否达到委托测试文件规定的要求，并形成可供委托方、测试机构及验收使用的客观测试结论。",
    ) or replace_first_paragraph_containing(
        doc,
        "测试目的是验证",
        "测试目的是验证智能磨损识别算法在规定测试图像集上对镜片磨损目标的识别结果，与人工确认并冻结的基准结果之间的一致程度是否达到委托测试文件规定的要求，并形成可供委托方、测试机构及验收使用的客观测试结论。",
    )
    replace_first_paragraph_containing(
        doc,
        "委托方负责提出测试需求，提供被测软件",
        "委托方负责提出测试需求，提交智能磨损识别算法的软件包或部署版本、运行说明、版本信息、必要配置文件、测试图像集、基准结果以及磨损镜片图像清单等测试资料。",
    ) or replace_first_paragraph_containing(
        doc,
        "委托方负责提出测试需求",
        "委托方负责提出测试需求，提交智能磨损识别算法的软件包或部署版本、运行说明、版本信息、必要配置文件、测试图像集、基准结果以及磨损镜片图像清单等测试资料。",
    )
    replace_first_paragraph_containing(
        doc,
        "测试机构负责依据本大纲及委托测试文件实施测试",
        "测试机构负责依据本大纲及委托测试要求，对智能磨损识别算法进行测试准备审查、环境核验、测试实施、结果核验、重复性核验和测试报告编制。",
    ) or replace_first_paragraph_containing(
        doc,
        "测试机构负责依据本大纲",
        "测试机构负责依据本大纲及委托测试要求，对智能磨损识别算法进行测试准备审查、环境核验、测试实施、结果核验、重复性核验和测试报告编制。",
    )
    replace_first_paragraph_containing(
        doc,
        "测试机构应确保以下信息完整记录并可追溯：被测软件名称及版本",
        "测试机构应确保以下信息完整记录并可追溯：智能磨损识别算法名称及版本、模型或算法版本、配置文件版本、测试图像集版本、基准结果版本、磨损镜片图像清单、运行环境信息、测试执行命令或运行方式、试运行结果、正式测试结果、重复性核验结果、原始输出结果文件、统计结果文件、测试日志以及异常情况记录。",
    ) or replace_first_paragraph_containing(
        doc,
        "测试机构应确保以下信息完整记录并可追溯",
        "测试机构应确保以下信息完整记录并可追溯：智能磨损识别算法名称及版本、模型或算法版本、配置文件版本、测试图像集版本、基准结果版本、磨损镜片图像清单、运行环境信息、测试执行命令或运行方式、试运行结果、正式测试结果、重复性核验结果、原始输出结果文件、统计结果文件、测试日志以及异常情况记录。",
    )

    # equations and inserted detailed sections
    set_paragraph_text(doc.paragraphs[56], "δ_i = 1（识别结果满足本次测试确认的一致性判定规则）", red=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph_text(doc.paragraphs[58], "A = (1/N) × Σ(i=1→N) δ_i × 100%", red=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # add references after 14
    current = doc.paragraphs[14]
    for text in [
        "《青少年近视防控功能镜片关键参数计量技术研究》项目任务书（项目编号：2024YFC2419500）",
        "《离焦微结构镜片磨损程度检测及关键参数评估研究》课题任务书（课题编号：2024YFC2419504）",
    ]:
        current = insert_paragraph_after(current, text, style="List Paragraph", red=True)

    # insert stronger test preparation details after 38
    current = doc.paragraphs[38]
    for text in [
        "智能磨损识别算法宜具备磨损镜片图像读取、识别处理、结果输出和运行记录等基本能力，以支持测试机构按照磨损镜片图像清单完成批量测试；",
        "测试前应形成磨损镜片图像准备清单。每一张磨损镜片图像及其对应基准结果应作为一个独立测试单元进行登记，清单中宜明确图像标识、数据版本、基准结果版本、运行命令、输出目录和记录责任人；",
        "磨损镜片图像应以测试图像及其对应基准结果成组组织，图像标识、图像文件、基准结果和统计结果之间应能够建立一一对应关系，以便后续测量、核验和追溯；",
        "磨损镜片图像准备内容宜包括：图像标识、图像文件路径、基准结果引用方式、识别结果输出形式、统计方式以及对应记录项。",
    ]:
        current = insert_paragraph_after(current, text, style="List Paragraph", red=True)

    # insert image examples after 37
    anchor = doc.paragraphs[37]
    current = insert_paragraph_after(
        anchor,
        "示例图像用于说明测试数据的基本形态。正式测试所采用的样本数据，以测试实施阶段确认的测试图像集为准。",
        style="Normal",
        red=True,
    )
    for idx, fig in enumerate(figs, start=1):
        p = insert_paragraph_after(current, style="Normal", red=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(fig), width=Cm(7.5))
        current = insert_paragraph_after(
            p,
            f"图 {idx} 磨损镜片图像示例",
            style="Normal",
            red=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    # insert data and environment detail after 42 and 52
    current = doc.paragraphs[42]
    for text in [
        "正式测试前，应确认测试图像集版本、基准结果版本、磨损镜片图像清单文件及其对应关系，并完成测试数据登记；",
        "测试过程中使用的统计规则、输出格式和结果留存方式，应在正式测试开始前完成确认；",
        "必要时可结合少量磨损镜片图像开展试运行，用于确认智能磨损识别算法运行状态和结果输出完整性；试运行结果不纳入正式统计。",
    ]:
        current = insert_paragraph_after(current, text, style="List Paragraph", red=True)

    current = doc.paragraphs[52]
    for text in [
        "正式测试与重复性核验宜采用同一套运行环境、同一份配置文件、同一组测试数据以及同一统计脚本；",
        "测试执行前应检查磨损镜片图像访问、结果输出路径、日志写入路径及必要权限是否满足要求，以避免因环境差异影响测试结果形成；",
        "对于支持批量处理的智能磨损识别算法，测试环境还应支持按磨损镜片图像清单顺序稳定运行，并保持图像输入、识别结果输出和日志记录之间的对应关系。",
    ]:
        current = insert_paragraph_after(current, text, style="List Paragraph", red=True)

    # insert repeatability criteria after 71
    current = doc.paragraphs[71]
    for text in [
        "每次正式测试应执行 1 次正式测试和 1 次重复性核验；",
        "正式测试结果与重复性核验结果的主指标绝对差值不大于 0.001 时，可认为测试结果具有良好重复性；",
        "若一致性判断不满足要求，应复核环境、数据、配置、输出文件和统计过程，并在查明原因后再形成正式结论。",
    ]:
        current = insert_paragraph_after(current, text, style="List Paragraph", red=True)

    # final terminology sweep for cross-version consistency
    replace_in_all_paragraphs(
        doc,
        "被测对象为能够接收磨损镜片图像、完成磨损目标识别并输出可统计结果的智能磨损识别算法。",
        "本次接受测试的内容为能够接收磨损镜片图像、完成磨损目标识别并输出可统计结果的智能磨损识别算法。",
    )
    replace_in_all_paragraphs(doc, "按样本标识输出对应结果", "按图像标识输出对应结果")
    replace_in_all_paragraphs(doc, "被测对象", "本次接受测试的内容")

    # appendix with one title and two logical tables
    current = doc.paragraphs[-1]
    current = insert_paragraph_after(current, "附录A 测试执行与重复性核验记录表（示例）", style="Heading 1", red=True)
    current = insert_paragraph_after(
        current,
        "附录A用于记录正式测试和重复性核验的关键条件、执行过程和结果信息。为便于填写和核查，附录A分为正式测试执行记录表和重复性核验记录表两部分。",
        style="Normal",
        red=True,
    )
    current = insert_paragraph_after(current, "A.1 正式测试执行记录表", style="Normal", red=True)
    table = insert_table_after(current, rows=17, cols=3)
    headers = ["分类", "记录项目", "记录内容"]
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, red=True, bold=True, center=True)
    items = [
        ("基本信息", "测试编号"),
        ("基本信息", "测试日期"),
        ("基本信息", "记录人员"),
        ("智能磨损识别算法信息", "算法名称及版本"),
        ("智能磨损识别算法信息", "模型/算法版本"),
        ("智能磨损识别算法信息", "配置文件版本"),
        ("测试数据与环境", "测试数据集版本"),
        ("测试数据与环境", "基准结果版本"),
        ("测试数据与环境", "磨损镜片图像清单路径"),
        ("测试数据与环境", "运行环境说明"),
        ("测试执行", "执行命令"),
        ("测试执行", "输入位置/输出目录"),
        ("测试执行", "日志文件路径"),
        ("测试执行", "开始时间 / 结束时间"),
        ("测试结果", "主指标结果"),
        ("测试结果", "结果文件完整性检查"),
    ]
    for r, (group, item) in enumerate(items, start=1):
        set_cell(table.rows[r].cells[0], group, red=True)
        set_cell(table.rows[r].cells[1], item, red=True)
        set_cell(table.rows[r].cells[2], " ", red=False)
    merge_specs = [(1, 3), (4, 6), (7, 10), (11, 14), (15, 16)]
    merged_labels = ["基本信息", "智能磨损识别算法信息", "测试数据与环境", "测试执行", "测试结果"]
    for (start, end), label in zip(merge_specs, merged_labels):
        top = table.cell(start, 0)
        for row in range(start + 1, end + 1):
            top = top.merge(table.cell(row, 0))
        set_cell(top, label, red=True, center=True)

    current = insert_paragraph_after_table(table, "A.2 重复性核验记录表", style="Normal", red=True)
    table2 = insert_table_after(current, rows=11, cols=3)
    for i, h in enumerate(headers):
        set_cell(table2.rows[0].cells[i], h, red=True, bold=True, center=True)
    items2 = [
        ("核验条件", "对应正式测试编号"),
        ("核验条件", "核验日期"),
        ("核验条件", "核验人员"),
        ("核验条件", "环境一致性核对"),
        ("核验条件", "数据与版本一致性核对"),
        ("核验执行", "核验执行命令"),
        ("核验执行", "日志文件路径"),
        ("核验结果", "重复性核验主指标结果"),
        ("核验结果", "与正式测试主指标差值"),
        ("核验结果", "一致性判断结论"),
    ]
    for r, (group, item) in enumerate(items2, start=1):
        set_cell(table2.rows[r].cells[0], group, red=True)
        set_cell(table2.rows[r].cells[1], item, red=True)
        set_cell(table2.rows[r].cells[2], " ", red=False)
    merge_specs2 = [(1, 5), (6, 7), (8, 10)]
    merged_labels2 = ["核验条件", "核验执行", "核验结果"]
    for (start, end), label in zip(merge_specs2, merged_labels2):
        top = table2.cell(start, 0)
        for row in range(start + 1, end + 1):
            top = top.merge(table2.cell(row, 0))
        set_cell(top, label, red=True, center=True)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    if args.sync_current:
        shutil.copyfile(output_path, output_path.parent / "CNAS测试大纲_当前版.docx")


if __name__ == "__main__":
    main()
